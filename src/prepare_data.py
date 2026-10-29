from pathlib import Path
import csv
import shutil
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "processed"
OUT.mkdir(parents=True, exist_ok=True)

# Convert MSOLVE CSVs to UTF-8 without modifying the originals
def convert_csv(source_path, output_path):
    with source_path.open("r", newline="", encoding="cp1252") as source_file:
        with output_path.open("w", newline="", encoding="utf-8-sig") as output_file:
            reader = csv.reader(source_file)
            writer = csv.writer(output_file)
            writer.writerows(reader)

convert_csv(RAW / "Msolve incident_Oct 26.csv", OUT / "msolve_incident.csv")
convert_csv(RAW / "Msolve service req_Oct 26.csv", OUT / "msolve_sr.csv")

# Export Alfa workbook sheets separately
workbook_path = RAW / "Alfa_Open_Resolved_Closed Tickets_Oct 26.xlsx"
workbook = load_workbook(workbook_path, read_only=True, data_only=True)

sheet_mapping = {
    "Open Tickets_Oct 26": "alfa_open.csv",
    "Resolved Tickets_Oct 26": "alfa_resolved.csv",
    "Closed Tickets_Oct 26": "alfa_closed.csv",
}

for sheet_name, filename in sheet_mapping.items():
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)

    output_path = OUT / filename
    with output_path.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(rows)

workbook.close()

# Report record counts and column counts
print("\nDATA PREPARATION REPORT")
print("-" * 55)

for path in sorted(OUT.glob("*.csv")):
    with path.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.reader(file)
        headers = next(reader, [])
        records = sum(1 for _ in reader)

    print(
        f"{path.name}: {records:,} records, "
        f"{len(headers)} columns"
    )

print(f"\nOutput directory: {OUT}")
